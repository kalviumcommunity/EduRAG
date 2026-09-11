import pymupdf as fitz
import docx
from dataclasses import dataclass, field
from typing import List, Optional, Dict, Any
from app.core.exceptions import ValidationException, AppException
from app.core.logging import logger


@dataclass
class ExtractedPage:
    page_number: Optional[int]
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)


class DocumentLoader:
    @staticmethod
    def load_pdf(file_path: str) -> List[ExtractedPage]:
        pages = []
        try:
            doc = fitz.open(file_path)
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text") or ""
                if text.strip():
                    pages.append(
                        ExtractedPage(
                            page_number=page_num + 1,  # 1-indexed page numbering
                            content=text,
                            metadata={"total_pages": len(doc)},
                        )
                    )
            doc.close()
        except Exception as e:
            logger.error(f"Failed to extract PDF content from {file_path}: {e}")
            raise AppException(f"PDF extraction failed: {str(e)}", error_code="EXTRACTION_FAILED")

        return pages

    @staticmethod
    def load_docx(file_path: str) -> List[ExtractedPage]:
        try:
            doc = docx.Document(file_path)
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            if not full_text.strip():
                return []
            # DOCX does not have native page boundaries, return single page document with page_number=None or 1
            return [
                ExtractedPage(
                    page_number=None,
                    content=full_text,
                    metadata={"paragraph_count": len(doc.paragraphs)},
                )
            ]
        except Exception as e:
            logger.error(f"Failed to extract DOCX content from {file_path}: {e}")
            raise AppException(f"DOCX extraction failed: {str(e)}", error_code="EXTRACTION_FAILED")

    @staticmethod
    def load_txt(file_path: str) -> List[ExtractedPage]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            if not content.strip():
                return []
            return [
                ExtractedPage(
                    page_number=None,
                    content=content,
                    metadata={"file_type": "txt"},
                )
            ]
        except Exception as e:
            logger.error(f"Failed to extract TXT content from {file_path}: {e}")
            raise AppException(f"TXT extraction failed: {str(e)}", error_code="EXTRACTION_FAILED")

    @classmethod
    def load_file(cls, file_path: str, file_extension: str) -> List[ExtractedPage]:
        ext = file_extension.lower().lstrip(".")
        if ext == "pdf":
            return cls.load_pdf(file_path)
        elif ext in ["docx", "doc"]:
            return cls.load_docx(file_path)
        elif ext == "txt":
            return cls.load_txt(file_path)
        else:
            raise ValidationException(f"Unsupported file format: .{ext}")
